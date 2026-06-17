import os
import json
import re
import zipfile
from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from fpdf import FPDF
from docx import Document

app = Flask(__name__)
CORS(app)

# --- CONFIGURATION ---
POEMS_DIR = 'markdown_poems'
MANIFEST_PATH = 'poems.json'
AUTHOR_NAME = "Urangani T.M"
INPUT_DOCX = 'Reality-a-colorful-illusion.docx'

# TAG_MAP for auto-tagging
TAG_MAP = {
    'nature': ['tree', 'flower', 'rain', 'moon', 'stars', 'lake', 'sun', 'sky', 'earth'],
    'love': ['love', 'heart', 'kiss', 'dear', 'princess', 'sister', 'her', 'she'],
    'darkness': ['hell', 'misery', 'deadly', 'karma', 'damned', 'suffer', 'pains'],
    'identity': ['who am i', 'reality', 'illusion', 'mind', 'dreamer', 'fool'],
    'time': ['eternity', 'beyond', 'new year', 'once upon a time', 'gone', 'forever']
}

# --- CORE UTILITIES ---

def parse_yaml(content):
    metadata = {}
    if content.startswith('---'):
        parts = content.split('---')
        if len(parts) >= 3:
            yaml_block = parts[1]
            lines = yaml_block.strip().split('\n')
            for line in lines:
                if ':' in line:
                    key, value = line.split(':', 1)
                    key = key.strip()
                    value = value.strip().strip('"').strip("'")
                    if value.startswith('[') and value.endswith(']'):
                        value = [t.strip().strip('"').strip("'") for t in value[1:-1].split(',') if t.strip()]
                    metadata[key] = value
    return metadata

def get_auto_tags(title, content):
    found_tags = set()
    text = (title + " " + content).lower()
    for tag, keywords in TAG_MAP.items():
        if any(kw in text for kw in keywords):
            found_tags.add(tag)
    return list(found_tags)

# --- ENGINE ACTIONS ---

def publish_manifest():
    files = sorted([f for f in os.listdir(POEMS_DIR) if f.endswith('.md')])
    poems_manifest = []
    for filename in files:
        filepath = os.path.join(POEMS_DIR, filename)
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
        metadata = parse_yaml(content)
        if metadata.get('type') == 'poem' and metadata.get('status') == 'published':
            poems_manifest.append({
                "title": metadata.get('title', 'Untitled'),
                "file": f"{POEMS_DIR}/{filename}",
                "tags": metadata.get('tags', []),
                "id": metadata.get('id')
            })
    with open(MANIFEST_PATH, 'w', encoding='utf-8') as f:
        json.dump(poems_manifest, f, indent=2)
    return len(poems_manifest)

def export_pdf():
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    
    files = sorted([f for f in os.listdir(POEMS_DIR) if f.endswith('.md')])
    for filename in files:
        filepath = os.path.join(POEMS_DIR, filename)
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
        metadata = parse_yaml(content)
        if metadata.get('status') != 'published':
            continue
            
        pdf.add_page()
        
        # Title
        pdf.set_font("Times", 'B', 24)
        pdf.cell(0, 40, metadata.get('title', 'Untitled'), ln=True, align='C')
        
        # Body
        pdf.set_font("Times", size=12)
        body = strip_yaml(content)
        # Simple Markdown parsing for PDF (Removing # Title and keeping line breaks)
        lines = body.split('\n')
        for line in lines:
            if line.startswith('# '): continue
            pdf.multi_cell(0, 10, line.strip(), align='C')
            
    output_path = "Reality_A_Colorful_Illusion.pdf"
    pdf.output(output_path)
    return output_path

def strip_yaml(text):
    if text.startswith('---'):
        parts = text.split('---')
        if len(parts) >= 3:
            return '---'.join(parts[2:]).strip()
    return text

# --- API ENDPOINTS ---

@app.route('/api/poems', methods=['GET'])
def get_poems():
    files = sorted([f for f in os.listdir(POEMS_DIR) if f.endswith('.md')])
    poems = []
    for filename in files:
        filepath = os.path.join(POEMS_DIR, filename)
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
        metadata = parse_yaml(content)
        poems.append({
            "filename": filename,
            "metadata": metadata,
            "content": strip_yaml(content)
        })
    return jsonify(poems)

@app.route('/api/poems', methods=['POST'])
def save_poem():
    data = request.json
    filename = data.get('filename')
    metadata = data.get('metadata')
    content = data.get('content')
    
    yaml_block = "---\n"
    for k, v in metadata.items():
        if isinstance(v, list):
            yaml_block += f"{k}: {json.dumps(v)}\n"
        else:
            yaml_block += f"{k}: \"{v}\"\n"
    yaml_block += "---\n\n"
    
    full_content = yaml_block + content
    filepath = os.path.join(POEMS_DIR, filename)
    
    with open(filepath, 'w', encoding='utf-8') as f:
        f.write(full_content)
    
    publish_manifest()
    return jsonify({"status": "success"})

@app.route('/api/export/pdf', methods=['GET'])
def download_pdf():
    path = export_pdf()
    return send_file(path, as_attachment=True)

@app.route('/api/export/zip', methods=['GET'])
def download_zip():
    zip_path = "poetry_collection.zip"
    with zipfile.ZipFile(zip_path, 'w') as zipf:
        for root, dirs, files in os.walk(POEMS_DIR):
            for file in files:
                zipf.write(os.path.join(root, file), file)
    return send_file(zip_path, as_attachment=True)

def import_from_word():
    if not os.path.exists(INPUT_DOCX):
        print(f"❌ Error: {INPUT_DOCX} not found.")
        return
    
    doc = Document(INPUT_DOCX)
    poems = []
    current_poem = None
    first_poem_found = False
    
    def is_num(text): return bool(re.match(r'^\s*\d+\.?\s*$', text))

    for para in doc.paragraphs:
        text = para.text.strip()
        if is_num(text):
            if current_poem: poems.append(current_poem)
            first_poem_found = True
            num = re.search(r'\d+', text).group()
            current_poem = {'number': num, 'title': None, 'lines': []}
            continue
        if not first_poem_found: continue
        if current_poem is not None:
            if current_poem['title'] is None and text: current_poem['title'] = text
            current_poem['lines'].append(para.text)
    
    if current_poem: poems.append(current_poem)

    new_count = 0
    for poem in poems:
        title = poem['title'] or "Untitled"
        filename = f"{poem['number'].zfill(3)}_{title.replace(' ', '_')[:30]}.md"
        path = os.path.join(POEMS_DIR, filename)
        if os.path.exists(path): continue
        
        with open(path, 'w', encoding='utf-8') as f:
            f.write(f"---\ntitle: \"{title}\"\nstatus: \"draft\"\ntype: \"poem\"\ntags: []\nid: {poem['number']}\n---\n\n")
            f.write(f"# {title}\n\n")
            f.write('\n'.join([f"{l}  " for l in poem['lines']]))
        new_count += 1
    print(f"✅ Imported {new_count} new poems as drafts.")

if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        cmd = sys.argv[1]
        if cmd == "publish":
            count = publish_manifest()
            print(f"✅ Published {count} poems.")
        elif cmd == "export-pdf":
            path = export_pdf()
            print(f"✅ Exported to {path}")
        elif cmd == "import":
            import_from_word()
        elif cmd == "studio":
            print("🚀 Poetry Studio starting on http://localhost:5000")
            app.run(port=5000, debug=True)
    else:
        print("Usage: python poetry_core.py [publish|export-pdf|studio]")
